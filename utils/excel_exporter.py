# =============================================================================
# Excel Report Generator
# =============================================================================
# Builds Excel workbooks in memory from SQLAlchemy query results and returns
# them as BytesIO buffers, so nothing is written to the local filesystem (which
# is ephemeral on Vercel).
#
# The generator takes *query results* rather than reaching for a global session,
# which keeps it independent of SQLite vs PostgreSQL and lets the admin area
# reuse the same code with a different query scope.
#
# PRIVACY / SECURITY: password hashes, session tokens, API keys and any other
# credential are never written to a worksheet.
# =============================================================================

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import Config
from utils.timezones import to_display

# -----------------------------------------------------------------------------
# Shared styling
# -----------------------------------------------------------------------------
HEADER_FONT = Font(bold=True, size=11, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
CELL_ALIGNMENT = Alignment(horizontal='left', vertical='center')
CENTER_ALIGNMENT = Alignment(horizontal='center', vertical='center')
THIN_BORDER = Border(
    left=Side(style='thin', color='D5D8DC'),
    right=Side(style='thin', color='D5D8DC'),
    top=Side(style='thin', color='D5D8DC'),
    bottom=Side(style='thin', color='D5D8DC'),
)

# Columns we refuse to export, whatever the caller passes in. Matched against the
# normalised header (lowercase, spaces -> underscores), so both "Session ID" and
# "session_token" are caught.
FORBIDDEN_HEADERS = {
    'password', 'password_hash', 'passwordhash', 'secret', 'secret_key',
    'api_key', 'token', 'session_secret', 'user_agent_secret',
    # Live-session capability tokens authorise frame submission. They must
    # never be written to a downloadable file.
    'session_id', 'session_token', 'sessionid', 'sessiontoken',
    'live_session_token',
}


# -----------------------------------------------------------------------------
# Helpers
# -----------------------------------------------------------------------------
def _fmt_dt(value, fmt):
    """
    Format a timestamp for a worksheet cell.

    The value is converted to the configured display timezone first (UTC
    storage -> Asia/Kolkata wall clock by default), so exported sheets show
    the same times as the web pages. Tolerates None and naive/aware mixes.
    """
    if value is None:
        return ''
    value = to_display(value)
    if isinstance(value, datetime):
        try:
            return value.strftime(fmt)
        except (ValueError, OSError):
            return value.isoformat()
    return str(value)


def _write_sheet(ws, headers, rows, center_columns=()):
    """
    Write a header row plus data rows, with consistent styling.

    Args:
        ws: openpyxl worksheet
        headers: list of column titles
        rows: iterable of row value tuples/lists
        center_columns: 1-based column numbers to centre-align
    """
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    for row_idx, values in enumerate(rows, 2):
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = (
                CENTER_ALIGNMENT if col_idx in center_columns else CELL_ALIGNMENT
            )


def _finish_sheet(ws):
    """Apply column widths, freeze the header row and enable autofilter."""
    for col_idx in range(1, ws.max_column + 1):
        widest = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    widest = max(widest, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(widest + 4, 40)

    ws.freeze_panes = 'A2'

    if ws.max_row > 1:
        ws.auto_filter.ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'


def _assert_no_secrets(headers):
    """Fail loudly if a caller tries to export a credential column."""
    for header in headers:
        key = str(header).strip().lower().replace(' ', '_')
        if key in FORBIDDEN_HEADERS:
            raise ValueError(f'Refusing to export sensitive column: {header!r}')


# -----------------------------------------------------------------------------
# Row builders - shared between the user and admin exports
# -----------------------------------------------------------------------------
def _detection_rows(detections, username_lookup=None):
    """
    Detection rows for a worksheet.

    ``session_id`` is deliberately NOT exported. It is the live-session
    capability token that authorises frame submission, so it must never leave
    the application in a spreadsheet.
    """
    rows = []
    for detection in detections:
        detected_at = detection.detected_at
        username = detection.user.username if getattr(detection, 'user', None) else (username_lookup or '')
        rows.append([
            detection.id,
            username,
            _fmt_dt(detected_at, Config.DATE_DISPLAY_FORMAT),
            _fmt_dt(detected_at, Config.TIME_DISPLAY_FORMAT),
            (detection.detection_type or '').capitalize(),
            detection.emotion or '',
            round(detection.confidence, 2) if detection.confidence is not None else 0,
            detection.face_count if detection.face_count is not None else 0,
            round(detection.processing_time, 3) if detection.processing_time is not None else '',
            _fmt_dt(detected_at, Config.TIMESTAMP_DISPLAY_FORMAT),
        ])
    return rows


def _session_rows(sessions, username_lookup=None):
    """Live-session rows. The session token is never included."""
    rows = []
    for session in sessions:
        started = session.started_at
        ended = session.ended_at
        username = session.user.username if getattr(session, 'user', None) else (username_lookup or '')
        rows.append([
            session.id,
            username,
            _fmt_dt(started, Config.DATE_DISPLAY_FORMAT),
            _fmt_dt(started, Config.TIME_DISPLAY_FORMAT),
            _fmt_dt(ended, Config.DATE_DISPLAY_FORMAT),
            _fmt_dt(ended, Config.TIME_DISPLAY_FORMAT),
            session.duration_seconds or 0,
            session.dominant_emotion or '',
            round(session.average_confidence, 2) if session.average_confidence is not None else 0,
            session.total_detections or 0,
            _fmt_dt(started, Config.TIMESTAMP_DISPLAY_FORMAT),
        ])
    return rows


def _user_rows(users):
    """User rows for the admin export. password_hash is intentionally excluded."""
    rows = []
    for user in users:
        rows.append([
            user.id,
            user.username,
            user.email,
            user.role or '',
            'Active' if user.is_active else 'Inactive',
            user.login_count or 0,
            _fmt_dt(user.created_at, Config.TIMESTAMP_DISPLAY_FORMAT),
            _fmt_dt(user.updated_at, Config.TIMESTAMP_DISPLAY_FORMAT),
            _fmt_dt(user.last_login_at, Config.TIMESTAMP_DISPLAY_FORMAT),
        ])
    return rows


def _activity_rows(activities):
    rows = []
    for activity in activities:
        rows.append([
            activity.id,
            activity.username,
            activity.activity_type or '',
            activity.ip_address or '',
            (activity.user_agent or '')[:120],
            activity.details or '',
            _fmt_dt(activity.created_at, Config.TIMESTAMP_DISPLAY_FORMAT),
        ])
    return rows


# Exported timestamps are rendered in the display timezone (Asia/Kolkata by
# default), so the header states that zone rather than UTC.
_TZ_LABEL = getattr(Config, 'DISPLAY_TIMEZONE_LABEL', 'IST')

DETECTION_HEADERS = [
    'ID', 'User', 'Date', 'Time', 'Detection Type', 'Emotion',
    'Confidence (%)', 'Face Count', 'Processing Time (s)',
    f'Timestamp ({_TZ_LABEL})',
]

SESSION_HEADERS = [
    'ID', 'User', 'Start Date', 'Start Time', 'End Date', 'End Time',
    'Duration (seconds)', 'Dominant Emotion', 'Average Confidence (%)',
    'Total Detections', f'Started At ({_TZ_LABEL})',
]

USER_HEADERS = [
    'ID', 'Username', 'Email', 'Role', 'Account Status', 'Login Count',
    f'Joined ({_TZ_LABEL})', f'Last Updated ({_TZ_LABEL})',
    f'Last Login ({_TZ_LABEL})',
]

ACTIVITY_HEADERS = [
    'ID', 'Username', 'Activity', 'IP Address', 'User Agent', 'Details',
    f'Timestamp ({_TZ_LABEL})',
]


# -----------------------------------------------------------------------------
# Public API
# -----------------------------------------------------------------------------
def build_user_workbook(detections, live_sessions, username=''):
    """
    Two-sheet workbook for a single user's own records.

    Returns a BytesIO positioned at 0, or None when there is no data.
    """
    if not detections and not live_sessions:
        return None

    wb = Workbook()

    ws_detections = wb.active
    ws_detections.title = 'Detection History'
    _assert_no_secrets(DETECTION_HEADERS)
    _write_sheet(
        ws_detections,
        DETECTION_HEADERS,
        _detection_rows(detections, username),
        center_columns=(1, 7, 8),
    )
    _finish_sheet(ws_detections)

    ws_sessions = wb.create_sheet('Live Sessions')
    _assert_no_secrets(SESSION_HEADERS)
    _write_sheet(
        ws_sessions,
        SESSION_HEADERS,
        _session_rows(live_sessions, username),
        center_columns=(1, 7, 9, 10),
    )
    _finish_sheet(ws_sessions)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_admin_workbook(users, detections, live_sessions, activities):
    """
    Four-sheet workbook for the admin area: Users, Detection History,
    Live Sessions, Activity Logs.

    Never includes password hashes.
    """
    wb = Workbook()

    ws_users = wb.active
    ws_users.title = 'Users'
    _assert_no_secrets(USER_HEADERS)
    _write_sheet(ws_users, USER_HEADERS, _user_rows(users), center_columns=(1, 4, 5, 6))
    _finish_sheet(ws_users)

    ws_detections = wb.create_sheet('Detection History')
    _assert_no_secrets(DETECTION_HEADERS)
    _write_sheet(
        ws_detections, DETECTION_HEADERS, _detection_rows(detections), center_columns=(1, 7, 8)
    )
    _finish_sheet(ws_detections)

    ws_sessions = wb.create_sheet('Live Sessions')
    _assert_no_secrets(SESSION_HEADERS)
    _write_sheet(
        ws_sessions, SESSION_HEADERS, _session_rows(live_sessions), center_columns=(1, 7, 9, 10)
    )
    _finish_sheet(ws_sessions)

    ws_activities = wb.create_sheet('Activity Logs')
    _assert_no_secrets(ACTIVITY_HEADERS)
    _write_sheet(
        ws_activities, ACTIVITY_HEADERS, _activity_rows(activities), center_columns=(1, 3)
    )
    _finish_sheet(ws_activities)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
